from pathlib import Path
from copy import deepcopy
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path.cwd();F=R/'submission_refined';d=Document(R/'.refine_build/raw.docx')
for p in d.paragraphs:
 if p.style.name=='Heading 2':p.style='Heading 1'
 elif p.style.name=='Heading 3':p.style='Heading 2'
def font(r,size=12,bold=None):
 r.font.name='Times New Roman';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0)
 if bold is not None:r.bold=bold
 r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Songti SC')
def setprop(e,name,attrs):
 z=OxmlElement(name)
 for k,v in attrs.items():z.set(qn(k),v)
 e.append(z);return z
sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
for k in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,k,Cm(2.5))
for name in ['Normal','Body Text','First Paragraph']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Songti SC');pf=st.paragraph_format;pf.line_spacing=1.5;pf.space_after=Pt(0);pf.first_line_indent=Cm(.847);pf.widow_control=True
for name,size in [('Title',22),('Heading 1',15),('Heading 2',13),('Heading 3',12)]:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Heiti SC');pf=st.paragraph_format;pf.space_before=Pt(10);pf.space_after=Pt(6);pf.first_line_indent=Cm(0);pf.keep_with_next=True;pf.line_spacing=1.3
p=d.paragraphs[0];p.text='皮肤病灶分割模型比较\n与自动提示级联实验';p.style='Title';p.alignment=A.CENTER;p.paragraph_format.space_before=Cm(4);p.paragraph_format.space_after=Cm(2)
for p in d.paragraphs[1:4]:p.alignment=A.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(20)
first=next(p for p in d.paragraphs if p.text=='一 摘要')
toc=first.insert_paragraph_before('目　录');toc.style='Heading 1';toc.alignment=A.CENTER;toc.paragraph_format.page_break_before=True
sections=[p for p in d.paragraphs if p.style.name=='Heading 1' and p.text!='目　录']
page_map=json.loads((R/'.refine_build/page_map.json').read_text()) if (R/'.refine_build/page_map.json').exists() else {}
for n,h in enumerate(sections):
 p=first.insert_paragraph_before();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(9);p.paragraph_format.line_spacing=1.5
 name=h.text;b=setprop(h._p,'w:bookmarkStart',{'w:id':str(200+n),'w:name':f'section{n}'});setprop(h._p,'w:bookmarkEnd',{'w:id':str(200+n)})
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),f'section{n}');rr=OxmlElement('w:r');t=OxmlElement('w:t');t.text=name;rr.append(t);link.append(rr);p._p.append(link)
 p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 p.add_run('\t'+str(page_map.get(name,'')))
 h.paragraph_format.page_break_before=(n==0 or name.startswith(('六 ','参考','附录')))
region='main';inrefs=False;eq=0
for p in d.paragraphs:
 txt=p.text;pf=p.paragraph_format;pf.widow_control=True
 if txt.startswith('附录'):region='appendix'
 if txt=='参考文献':inrefs=True
 if txt.startswith('附录'):inrefs=False
 if p._p.findall('.//'+qn('w:drawing')):
  p.alignment=A.CENTER;pf.first_line_indent=Cm(0);pf.keep_with_next=True;pf.space_before=Pt(7);pf.space_after=Pt(2)
  maxh=12.5 if region=='main' else 17
  for ex in p._p.findall('.//'+qn('wp:extent')):
   ow,oh=int(ex.get('cx')),int(ex.get('cy'));w=int(Cm(16));h=int(oh*w/ow)
   if h>Cm(maxh):w=int(w*Cm(maxh)/h);h=int(Cm(maxh))
   ex.set('cx',str(w));ex.set('cy',str(h))
   for z in p._p.findall('.//'+qn('a:xfrm')+'/'+qn('a:ext')):z.set('cx',str(w));z.set('cy',str(h))
 elif re.match(r'^[图表]\d+　',txt):
  p.alignment=A.CENTER;pf.first_line_indent=Cm(0);pf.keep_with_next=True;pf.space_before=Pt(3);pf.space_after=Pt(3);pf.line_spacing=1.2
  for r in p.runs:font(r,11)
 elif txt.startswith('注：'):
  pf.first_line_indent=Cm(0);pf.space_after=Pt(6);pf.line_spacing=1.15
  for r in p.runs:font(r,9)
 elif inrefs and re.match(r'^\[\d+\]',txt):
  p.alignment=A.LEFT;pf.left_indent=Cm(.7);pf.first_line_indent=Cm(-.7);pf.line_spacing=1.3;pf.space_after=Pt(4);pf.keep_together=True
  for r in p.runs:font(r,10.5)
 if p._p.findall('.//'+qn('m:oMathPara')):
  eq+=1;pf.first_line_indent=Cm(0);pf.space_before=Pt(6);pf.space_after=Pt(6);p.alignment=A.CENTER
  # Explicit ordinary-text equation number survives Word and PDF export.
  for z in p._p.findall('.//'+qn('m:r')):
   ts=z.findall(qn('m:t'))
  p.add_run('　（'+str(eq)+'）')
 if txt.startswith('个人观点'):
  for r in p.runs:
   if r.bold:setprop(r._element.get_or_add_rPr(),'w:highlight',{'w:val':'yellow'})
 for h in p._p.findall('.//'+qn('w:hyperlink')):
  if (h.get(qn('w:anchor')) or '').startswith('ref-'):
   # Trim the actual preceding text node; no visible gap before superscript.
   prev=h.getprevious()
   if prev is not None:
    nodes=prev.findall('.//'+qn('w:t'))
    if nodes:nodes[-1].text=(nodes[-1].text or '').rstrip(' \u2060')+'\u2060'
   for rr in h.findall(qn('w:r')):
    pr=rr.find(qn('w:rPr'))
    if pr is None:pr=OxmlElement('w:rPr');rr.insert(0,pr)
    setprop(pr,'w:vertAlign',{'w:val':'superscript'});setprop(pr,'w:color',{'w:val':'000000'})
for tb in d.tables:
 tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False;n=len(tb.columns);wide=n>=7;avail=16
 widths=[avail/n]*n
 if n==8:widths=[3.0]+[13/7]*7
 elif n==7:widths=[3.6]+[12.4/6]*6
 elif n==3:widths=[4.5,5.75,5.75]
 elif n==4:widths=[4,4,4,4]
 elif n==6:widths=[3.8]+[12.2/5]*5
 for c,w in zip(tb.columns,widths):c.width=Cm(w)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','bottom','left','right','insideH','insideV']:setprop(borders,'w:'+side,{'w:val':'single','w:sz':'4','w:color':'BAC8D2'})
 tb._tbl.tblPr.append(borders)
 for i,row in enumerate(tb.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tc=c._tc.get_or_add_tcPr();setprop(tc,'w:shd',{'w:fill':'EAF2F8' if i==0 else 'FFFFFF'});mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:setprop(mar,'w:'+side,{'w:w':'55','w:type':'dxa'})
   tc.append(mar)
   for p in c.paragraphs:
    p.alignment=A.CENTER;pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.line_spacing=1.1;pf.space_before=Pt(1);pf.space_after=Pt(1)
    for r in p.runs:font(r,9 if wide else 10,i==0)
fp=sec.footer.paragraphs[0];fp.alignment=A.CENTER;fp.paragraph_format.first_line_indent=Cm(0);setprop(fp._p,'w:fldSimple',{'w:instr':'PAGE'})
sec.different_first_page_header_footer=True
for p in sec.header.paragraphs:p.text=''
d.core_properties.author='王春博';d.core_properties.title='皮肤病灶分割模型比较与自动提示级联实验'
out=F/'202620000613_王春博_机器学习课程作业_最终报告.docx';d.save(out);print(out,'equations',eq,'figures',len(d.inline_shapes))
